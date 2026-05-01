"""
RAG workflow orchestration utilities.

This module mirrors the workflow in RAG_workflow.png:
1. Pre-production: extract -> chunk -> embed -> store
2. Production retriever: embed query -> vector search -> top-k chunks
3. Production reader: build context -> prompt LLM -> answer
"""

from __future__ import annotations

import io
import logging
import uuid
from collections.abc import Generator
from dataclasses import asdict, dataclass
from pathlib import Path
from urllib.parse import quote

from docx import Document
from pypdf import PdfReader

from llm import generate_answer, generate_answer_stream
from vector_store import (
    add_conversation_turn,
    add_documents,
    get_conversation_history,
    query_similar,
)

logger = logging.getLogger(__name__)


ALLOWED_DOC_TYPES = {
    "application/pdf",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class WorkflowInputError(ValueError):
    """Raised when the input cannot be processed into a valid RAG stage."""


@dataclass(frozen=True)
class IngestionResult:
    filename: str
    upload_id: str
    source_url: str
    total_words: int
    total_chunks: int
    chunk_ids: list[str]
    tokens_per_chunk: int
    overlap: int
    chunks_preview: list[str]

    def to_dict(self) -> dict:
        return asdict(self)


def normalize_filename(filename: str | None) -> str:
    """Normalize filename to a safe leaf name only."""
    return Path(filename or "unknown").name


def doc_storage_path(storage_dir: Path, upload_id: str, filename: str) -> Path:
    """Get stable on-disk path for a stored source document."""
    return storage_dir / f"{upload_id}__{normalize_filename(filename)}"


def build_source_url(metadata: dict | None) -> str | None:
    """Build a browseable URL for a chunk's source document."""
    if not metadata:
        return None

    filename = metadata.get("filename")
    upload_id = metadata.get("upload_id")
    if not filename or not upload_id:
        return None

    return f"/weather/source-doc/{upload_id}/{quote(str(filename))}"


def format_sources(results: list[dict]) -> list[dict]:
    """Project vector results into API source objects with browse URLs."""
    sources = []
    for match in results:
        metadata = match.get("metadata") or {}
        # Keep source payload minimal and UI-friendly.
        sources.append(
            {
                "filename": metadata.get("filename"),
                "chunk": metadata.get("chunk_index"),
                "upload_id": metadata.get("upload_id"),
                "url": build_source_url(metadata),
            }
        )
    return sources


def extract_text_from_file(file_content: bytes, content_type: str) -> str:
    """Extract raw text from uploaded PDF or Word documents."""
    logger.info("Extracting text from file type: %s", content_type)

    if content_type == "application/pdf":
        reader = PdfReader(io.BytesIO(file_content))
        text = "".join([page.extract_text() or "" for page in reader.pages])
        logger.info("Extracted text from %d PDF pages", len(reader.pages))
        return text

    if content_type in {
        "application/msword",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }:
        doc = Document(io.BytesIO(file_content))
        text = "\n".join([para.text for para in doc.paragraphs])
        logger.info("Extracted text from %d paragraphs", len(doc.paragraphs))
        return text

    return ""


def chunk_text_by_tokens(
    text: str,
    tokens_per_chunk: int = 100,
    overlap: int = 20,
) -> list[str]:
    """
    Split text into overlapping chunks.

    Uses words as a lightweight token proxy.
    """
    words = text.split()
    if not words:
        return []

    if overlap >= tokens_per_chunk:
        logger.warning(
            "Overlap (%d) >= chunk size (%d). Adjusting overlap.",
            overlap,
            tokens_per_chunk,
        )
        overlap = max(0, tokens_per_chunk - 1)

    step = max(1, tokens_per_chunk - overlap)
    chunks: list[str] = []

    for i in range(0, len(words), step):
        chunk = words[i : i + tokens_per_chunk]
        if not chunk:
            break
        chunks.append(" ".join(chunk))
        if i + tokens_per_chunk >= len(words):
            break

    return chunks


def ingest_document(
    *,
    file_content: bytes,
    content_type: str,
    filename: str | None,
    storage_dir: Path,
    tokens_per_chunk: int = 100,
    overlap: int = 20,
) -> IngestionResult:
    """
    Pre-production workflow:
    upload -> extract -> chunk -> embed -> store.
    """
    safe_filename = normalize_filename(filename)

    if content_type not in ALLOWED_DOC_TYPES:
        raise WorkflowInputError("Only PDF and DOC/DOCX files are allowed.")

    text = extract_text_from_file(file_content, content_type)
    if not text.strip():
        raise WorkflowInputError("Could not extract text from the document.")

    chunks = chunk_text_by_tokens(text, tokens_per_chunk=tokens_per_chunk, overlap=overlap)
    if not chunks:
        raise WorkflowInputError("No chunks generated - document may be too short.")

    upload_id = str(uuid.uuid4())
    chunk_ids = add_documents(chunks, safe_filename, upload_id=upload_id)

    storage_dir.mkdir(parents=True, exist_ok=True)
    storage_path = doc_storage_path(storage_dir, upload_id, safe_filename)
    storage_path.write_bytes(file_content)

    return IngestionResult(
        filename=safe_filename,
        upload_id=upload_id,
        source_url=f"/weather/source-doc/{upload_id}/{quote(safe_filename)}",
        total_words=len(text.split()),
        total_chunks=len(chunks),
        chunk_ids=chunk_ids,
        tokens_per_chunk=tokens_per_chunk,
        overlap=overlap,
        chunks_preview=chunks[:3],
    )


def retrieve_top_k(query: str, n_results: int = 3) -> list[dict]:
    """Production retriever workflow: embed query -> vector search -> top-k."""
    results = query_similar(query, n_results=n_results)
    for result in results:
        # Attach a direct source-document URL for each retrieved chunk.
        result["source_url"] = build_source_url(result.get("metadata"))
    return results


def answer_query(
    *,
    query: str,
    n_results: int = 3,
    user_id: str = "anonymous",
    max_memory_turns: int = 5,
    model: str | None = None,
) -> dict:
    """Production reader workflow: retrieve -> build context -> generate answer."""
    results = retrieve_top_k(query, n_results=n_results)
    if not results:
        return {
            "query": query,
            "user_id": user_id,
            "answer": "No relevant context found to answer the question.",
            "sources": [],
        }

    user_history = get_conversation_history(user_id, max_turns=max_memory_turns)
    answer = generate_answer(query, results, history=user_history, model=model)
    add_conversation_turn(user_id, query, answer, max_turns=max_memory_turns)

    return {
        "query": query,
        "user_id": user_id,
        "answer": answer,
        "sources": format_sources(results),
    }


def stream_answer_query(
    *,
    query: str,
    n_results: int = 3,
    user_id: str = "anonymous",
    max_memory_turns: int = 5,
    model: str | None = None,
) -> tuple[Generator[str, None, None] | None, list[dict]]:
    """Streaming reader workflow with conversation persistence."""
    results = retrieve_top_k(query, n_results=n_results)
    if not results:
        return None, []

    user_history = get_conversation_history(user_id, max_turns=max_memory_turns)
    sources = format_sources(results)

    def token_stream() -> Generator[str, None, None]:
        answer_parts: list[str] = []
        for token in generate_answer_stream(query, results, history=user_history, model=model):
            answer_parts.append(token)
            yield token

        full_answer = "".join(answer_parts)
        # Persist only the completed answer once streaming finishes.
        add_conversation_turn(user_id, query, full_answer, max_turns=max_memory_turns)

    return token_stream(), sources
